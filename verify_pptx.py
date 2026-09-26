from pptx import Presentation
prs = Presentation('Mid_Defense_FINAL.pptx')
slides = list(prs.slides)
print(f'Total slides: {len(slides)}')
for i, slide in enumerate(slides):
    texts = []
    for shape in slide.shapes:
        if hasattr(shape, 'text') and shape.text.strip():
            t = shape.text.strip()[:70].replace('\n', ' ')
            texts.append(t)
    combined = ' || '.join(texts[:3])
    print(f'  Slide {i+1}: {combined}')
