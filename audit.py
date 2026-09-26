from pptx import Presentation
prs = Presentation('Mid_Defense_FINAL.pptx')
for i, sl in enumerate(prs.slides):
    print(f"\n=== Slide {i+1} ===")
    for sh in sl.shapes:
        if hasattr(sh, 'text') and sh.text.strip():
            print(f"  [{sh.name}] -> {sh.text.strip()[:80].replace(chr(10),' | ')}")
